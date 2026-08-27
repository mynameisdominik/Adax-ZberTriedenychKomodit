from io import BytesIO

from docx import Document

from adax_calendar import parse_docx


def test_parse_adax_table() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Mesiac"
    table.cell(0, 1).text = "Plast, VKM, Kov ob."
    table.cell(0, 2).text = "Papier"
    table.cell(1, 0).text = "Január"
    table.cell(1, 1).text = "29.01.2026"
    table.cell(1, 2).text = "29.01.2026"
    table.cell(2, 0).text = "Február"
    table.cell(2, 1).text = "26.02.2026"
    table.cell(2, 2).text = "26.02.2026"

    output = parse_docx(BytesIO(_document_bytes(document)))

    assert [item.as_dict() for item in output] == [
        {"date": "2026-01-29", "commodity": "Plast, VKM, Kov ob."},
        {"date": "2026-01-29", "commodity": "Papier"},
        {"date": "2026-02-26", "commodity": "Plast, VKM, Kov ob."},
        {"date": "2026-02-26", "commodity": "Papier"},
    ]


def _document_bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
